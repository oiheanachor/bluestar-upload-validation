"""
report_writer.py

Writes one Excel workbook containing Item and BOM validation summaries,
detailed results, protected BOM results, and the overall release decision.
"""

from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)
FAIL_FILL = PatternFill("solid", fgColor="F4CCCC")
WARNING_FILL = PatternFill("solid", fgColor="FFF2CC")
PASS_FILL = PatternFill("solid", fgColor="D9EAD3")


def count_failures(dataframe: pd.DataFrame) -> int:
    """Count rows whose Overall Result is FAIL."""
    if dataframe.empty or "Overall Result" not in dataframe.columns:
        return 0
    return int(dataframe["Overall Result"].astype(str).eq("FAIL").sum())


def overall_decision(
    item_results_df: pd.DataFrame,
    bom_results_df: pd.DataFrame,
    protected_bom_results_df: pd.DataFrame,
) -> str:
    """Return the release decision from all blocking validation results."""
    failure_count = (
        count_failures(item_results_df)
        + count_failures(bom_results_df)
        + count_failures(protected_bom_results_df)
    )
    return "NOT_READY_TO_RELEASE" if failure_count else "READY_TO_RELEASE"


def build_overview(
    item_results_df: pd.DataFrame,
    bom_results_df: pd.DataFrame,
    protected_bom_results_df: pd.DataFrame,
    environment: str,
    uploader: str,
    upload_start: Any,
    upload_end: Any,
) -> pd.DataFrame:
    """Build concise run metadata and overall control totals."""
    item_failures = count_failures(item_results_df)
    bom_failures = count_failures(bom_results_df)
    protected_failures = count_failures(protected_bom_results_df)

    return pd.DataFrame(
        {
            "Metric": [
                "Release Decision",
                "Environment",
                "Uploader",
                "Upload Start",
                "Upload End",
                "Report Generated UTC",
                "Item Result Rows",
                "Item Failure Rows",
                "BOM Result Rows",
                "BOM Failure Rows",
                "Protected BOMs Checked",
                "Protected BOM Failures",
            ],
            "Value": [
                overall_decision(
                    item_results_df,
                    bom_results_df,
                    protected_bom_results_df,
                ),
                environment,
                uploader,
                str(upload_start),
                str(upload_end) if upload_end is not None else "POST snapshot time",
                datetime.now(timezone.utc).isoformat(timespec="seconds"),
                len(item_results_df),
                item_failures,
                len(bom_results_df),
                bom_failures,
                len(protected_bom_results_df),
                protected_failures,
            ],
        }
    )


def safe_dataframe(dataframe: pd.DataFrame | None) -> pd.DataFrame:
    """Return a writable DataFrame for empty or optional results."""
    if dataframe is None:
        return pd.DataFrame({"Information": ["No results returned"]})
    if dataframe.empty and len(dataframe.columns) == 0:
        return pd.DataFrame({"Information": ["No results returned"]})
    return dataframe.copy()


def format_worksheet(worksheet) -> None:
    """Apply readable, restrained formatting to one worksheet."""
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    for cell in worksheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for column_index, cells in enumerate(worksheet.columns, start=1):
        values = ["" if cell.value is None else str(cell.value) for cell in cells]
        width = min(max(len(value) for value in values) + 2, 60)
        worksheet.column_dimensions[get_column_letter(column_index)].width = max(width, 10)

    result_column = None
    for cell in worksheet[1]:
        if cell.value == "Overall Result":
            result_column = cell.column
            break

    if result_column is not None:
        for row_index in range(2, worksheet.max_row + 1):
            result_cell = worksheet.cell(row=row_index, column=result_column)
            result_value = str(result_cell.value or "").upper()
            if result_value == "FAIL":
                result_cell.fill = FAIL_FILL
            elif result_value == "PASS_WITH_WARNING":
                result_cell.fill = WARNING_FILL
            elif result_value == "PASS":
                result_cell.fill = PASS_FILL


def write_validation_report(
    item_summary_df: pd.DataFrame,
    item_results_df: pd.DataFrame,
    bom_summary_df: pd.DataFrame,
    bom_results_df: pd.DataFrame,
    protected_bom_results_df: pd.DataFrame,
    environment: str,
    uploader: str,
    upload_start: Any,
    upload_end: Any = None,
    output_path: str | Path = "output/ValidationReport.xlsx",
) -> Path:
    """Write the complete validation workbook and return its path."""
    output_file = Path(output_path)
    output_file.parent.mkdir(parents=True, exist_ok=True)

    overview_df = build_overview(
        item_results_df=item_results_df,
        bom_results_df=bom_results_df,
        protected_bom_results_df=protected_bom_results_df,
        environment=environment,
        uploader=uploader,
        upload_start=upload_start,
        upload_end=upload_end,
    )

    worksheets = {
        "Overview": overview_df,
        "Item Summary": safe_dataframe(item_summary_df),
        "Item Results": safe_dataframe(item_results_df),
        "BOM Summary": safe_dataframe(bom_summary_df),
        "BOM Results": safe_dataframe(bom_results_df),
        "Protected BOMs": safe_dataframe(protected_bom_results_df),
    }

    with pd.ExcelWriter(output_file, engine="openpyxl") as writer:
        for sheet_name, dataframe in worksheets.items():
            dataframe.to_excel(writer, sheet_name=sheet_name, index=False)
            format_worksheet(writer.book[sheet_name])

    return output_file
